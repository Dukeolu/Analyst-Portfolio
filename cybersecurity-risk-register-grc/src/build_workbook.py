"""
Builds reports/risk-register.xlsx — an actual working Excel risk register,
not a static export: likelihood x impact scores and risk reduction are live
formulas, and the GRC Summary sheet rolls up the register with SUMIF/
COUNTIF rather than hardcoded numbers, so editing a risk's likelihood or
status on the Risk Register sheet updates the summary automatically.
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.utils import get_column_letter

with open("data/raw/risk_register.csv") as f:
    risks = list(csv.DictReader(f))

wb = Workbook()

# ---------------------------------------------------------------- styles
HEADER_FILL = PatternFill("solid", fgColor="8B2F2F")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10.5, name="Calibri")
TITLE_FONT = Font(bold=True, size=14, name="Calibri", color="8B2F2F")
SUBTITLE_FONT = Font(italic=True, size=9.5, color="524D4A", name="Calibri")
BODY_FONT = Font(size=10, name="Calibri")
THIN = Side(style="thin", color="C2B8B2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")

# ---------------------------------------------------------- Risk Register
ws = wb.active
ws.title = "Risk Register"
ws["A1"] = "Meridian Manufacturing (simulated) — Cybersecurity & IT Risk Register"
ws["A1"].font = TITLE_FONT
ws["A2"] = "Rolled up from every case in the IT & Cybersecurity portfolio track. Scores are live formulas — edit likelihood, impact, or residual values below and the numbers recalculate."
ws["A2"].font = SUBTITLE_FONT
ws.merge_cells("A1:H1")
ws.merge_cells("A2:P2")

headers = [
    "Risk ID", "Title", "Category", "CSF Function", "Source Case",
    "Inh. Likelihood", "Inh. Impact", "Inherent Score",
    "Res. Likelihood", "Res. Impact", "Residual Score",
    "Risk Reduction %", "Status", "Owner", "Target Date", "Notes",
]
HEADER_ROW = 4
for col, h in enumerate(headers, start=1):
    c = ws.cell(row=HEADER_ROW, column=col, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER

status_fills = {
    "Closed": PatternFill("solid", fgColor="D7EFD7"),
    "Open": PatternFill("solid", fgColor="F5D6D6"),
    "In Progress": PatternFill("solid", fgColor="FBEAC7"),
    "Risk Accepted": PatternFill("solid", fgColor="E3DCF5"),
    "Monitoring": PatternFill("solid", fgColor="DCEAF5"),
}

start_row = HEADER_ROW + 1
for i, r in enumerate(risks):
    row = start_row + i
    ws.cell(row=row, column=1, value=r["risk_id"])
    ws.cell(row=row, column=2, value=r["title"]).alignment = WRAP
    ws.cell(row=row, column=3, value=r["category"])
    ws.cell(row=row, column=4, value=r["csf_function"])
    ws.cell(row=row, column=5, value=r["source_case"]).alignment = WRAP
    ws.cell(row=row, column=6, value=int(r["inherent_likelihood"]))
    ws.cell(row=row, column=7, value=int(r["inherent_impact"]))
    # live formula: inherent score = likelihood x impact
    ws.cell(row=row, column=8, value=f"=F{row}*G{row}")
    ws.cell(row=row, column=9, value=int(r["residual_likelihood"]))
    ws.cell(row=row, column=10, value=int(r["residual_impact"]))
    # live formula: residual score = likelihood x impact
    ws.cell(row=row, column=11, value=f"=I{row}*J{row}")
    # live formula: risk reduction % = (inherent - residual) / inherent
    ws.cell(row=row, column=12, value=f"=IF(H{row}=0,0,ROUND((H{row}-K{row})/H{row}*100,1))")
    status_cell = ws.cell(row=row, column=13, value=r["status"])
    status_cell.fill = status_fills.get(r["status"], PatternFill())
    ws.cell(row=row, column=14, value=r["owner"])
    ws.cell(row=row, column=15, value=r["target_date"])
    ws.cell(row=row, column=16, value=r["notes"]).alignment = WRAP
    for col in range(1, 17):
        ws.cell(row=row, column=col).border = BORDER
        if col not in (2, 5, 16):
            ws.cell(row=row, column=col).font = BODY_FONT
    ws.row_dimensions[row].height = 34

last_row = start_row + len(risks) - 1

widths = [8, 34, 14, 12, 26, 9, 8, 10, 9, 8, 10, 11, 13, 15, 12, 40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = f"A{start_row}"

# conditional color scale on residual score (green=low risk, red=high risk)
ws.conditional_formatting.add(
    f"K{start_row}:K{last_row}",
    ColorScaleRule(
        start_type="min", start_color="63BE7B",
        mid_type="percentile", mid_value=50, mid_color="FFEB84",
        end_type="max", end_color="C0392B",
    ),
)

# ---------------------------------------------------------------- GRC Summary
ws2 = wb.create_sheet("GRC Summary")
ws2["A1"] = "GRC Summary — live rollup of the Risk Register sheet"
ws2["A1"].font = TITLE_FONT
ws2["A2"] = "Every value below is a formula referencing 'Risk Register' — no hand-typed totals."
ws2["A2"].font = SUBTITLE_FONT
ws2.merge_cells("A1:D1")
ws2.merge_cells("A2:D2")

reg = "'Risk Register'"
summary_rows = [
    ("Total risks tracked", f"=COUNTA({reg}!A{start_row}:A{last_row})"),
    ("Total inherent risk score", f"=SUM({reg}!H{start_row}:H{last_row})"),
    ("Total residual risk score", f"=SUM({reg}!K{start_row}:K{last_row})"),
    ("Overall risk reduction %",
     f"=ROUND((B{6}-B{7})/B{6}*100,1)"),  # filled below once row numbers known
    ("", ""),
    ("Status breakdown", ""),
    ("Closed", f"=COUNTIF({reg}!M{start_row}:M{last_row},\"Closed\")"),
    ("Open", f"=COUNTIF({reg}!M{start_row}:M{last_row},\"Open\")"),
    ("In Progress", f"=COUNTIF({reg}!M{start_row}:M{last_row},\"In Progress\")"),
    ("Risk Accepted", f"=COUNTIF({reg}!M{start_row}:M{last_row},\"Risk Accepted\")"),
    ("Monitoring", f"=COUNTIF({reg}!M{start_row}:M{last_row},\"Monitoring\")"),
    ("", ""),
    ("Residual risk by CSF function", ""),
    ("Govern", f"=SUMIF({reg}!D{start_row}:D{last_row},\"Govern\",{reg}!K{start_row}:K{last_row})"),
    ("Identify", f"=SUMIF({reg}!D{start_row}:D{last_row},\"Identify\",{reg}!K{start_row}:K{last_row})"),
    ("Protect", f"=SUMIF({reg}!D{start_row}:D{last_row},\"Protect\",{reg}!K{start_row}:K{last_row})"),
    ("Detect", f"=SUMIF({reg}!D{start_row}:D{last_row},\"Detect\",{reg}!K{start_row}:K{last_row})"),
    ("Respond", f"=SUMIF({reg}!D{start_row}:D{last_row},\"Respond\",{reg}!K{start_row}:K{last_row})"),
    ("Recover", f"=SUMIF({reg}!D{start_row}:D{last_row},\"Recover\",{reg}!K{start_row}:K{last_row})"),
]

row0 = 4
for i, (label, formula) in enumerate(summary_rows):
    row = row0 + i
    lc = ws2.cell(row=row, column=1, value=label)
    if label in ("Status breakdown", "Residual risk by CSF function"):
        lc.font = Font(bold=True, size=11, name="Calibri", color="6B2020")
    else:
        lc.font = BODY_FONT
    if formula:
        vc = ws2.cell(row=row, column=2, value=formula)
        vc.font = Font(bold=True, size=10.5, name="Calibri")

# fix the overall-reduction formula to point at the correct rows (inherent=row0+1, residual=row0+2)
ws2[f"B{row0+3}"] = f"=ROUND((B{row0+1}-B{row0+2})/B{row0+1}*100,1)"

ws2.column_dimensions["A"].width = 32
ws2.column_dimensions["B"].width = 14

wb.save("reports/risk-register.xlsx")
print("Wrote reports/risk-register.xlsx")
